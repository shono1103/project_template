"""**既に貼ってある**資料の画像を「土台 (結合セル) + 実寸アンカー」に組み替える。

    python fix_anchors.py <in.xlsx> <out.xlsx>

`relayout.py` + `place_images.py` で作り直せないとき (= 資料が既に配られていて、
セルの中身や行番号を動かせないとき) の**修復用**。新規に作るなら使わない。

やること:

* 画像を `oneCellAnchor` + 明示 `<ext>` にする。**大きさが行高から独立するので歪まない。**
  縦横比は画像の実寸から取り直す。
* 画像が入る結合セルを **18.0pt の整数行だけ**で組み直す (= 土台)。
  端数行 (0.38pt 等) があると、行高に下限を持つ描画系 (Google スプレッドシート等) で
  土台が縮み、画像がラベル行にはみ出す。
* 継ぎ (直前の画像の真下に隙間無く続く分割画像) は鎖ごとに 1 つの土台へまとめ、
  2 枚目以降は `rowOff` で真下に詰める。**継ぎ目に白い隙間が出ない。**

**行数・行番号・セルの中身・ハイパーリンク・画像のバイト列は変えない。**
土台の行数が足りない画像があれば `assert` で止まる (その場合は器から作り直す)。
"""
import sys
from openpyxl import load_workbook
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils import get_column_letter

EMU, ROW_PX, ROW_PT = 9525, 24, 18.0
ROW_EMU = ROW_PX * EMU

if len(sys.argv) != 3:
    sys.exit(__doc__)
src, dst = sys.argv[1], sys.argv[2]
wb = load_workbook(src)
total = 0

for ws in wb.worksheets:
    if not ws._images:
        continue
    if any(getattr(im.anchor, "to", None) is None for im in ws._images):
        sys.exit("%s は既に oneCellAnchor になっている (組み替え済み)。二重に掛けない" % ws.title)
    print("===", ws.title)
    # 現在のアンカーから (画像, 上端行, 下端行, 左端列, 右端列) を作る
    box = []
    for im in sorted(ws._images, key=lambda i: i.anchor._from.row):
        a = im.anchor
        box.append((im, a._from.row + 1, a.to.row, a._from.col + 1, a.to.col))

    # 真下に続くものを 1 つの鎖にまとめる
    chains, cur = [], [box[0]]
    for b in box[1:]:
        if b[1] == cur[-1][2] + 1:
            cur.append(b)
        else:
            chains.append(cur)
            cur = [b]
    chains.append(cur)

    for ch in chains:
        R0, R1, c0, c1 = ch[0][1], ch[-1][2], ch[0][3], ch[0][4]
        w_px = sum(round(ws.column_dimensions[get_column_letter(c)].width * 7) + 5
                   for c in range(c0, c1 + 1))
        cx = w_px * EMU

        for r in range(R0, R1 + 1):                     # 土台は 18.0pt の整数行だけで組む
            ws.row_dimensions[r].height = ROW_PT
        for rng in list(ws.merged_cells.ranges):        # 鎖の分の結合を 1 つにまとめ直す
            if (rng.min_col == c0 and rng.max_col == c1
                    and R0 <= rng.min_row and rng.max_row <= R1):
                ws.unmerge_cells(str(rng))
        ws.merge_cells(start_row=R0, start_column=c0, end_row=R1, end_column=c1)

        off = 0
        for im, _r0, _r1, _c0, _c1 in ch:
            cy = round(cx * im.height / im.width)       # 縦横比は画像そのものから取る
            im.anchor = OneCellAnchor(
                _from=AnchorMarker(col=c0 - 1, colOff=0,
                                   row=R0 - 1 + off // ROW_EMU, rowOff=off % ROW_EMU),
                ext=XDRPositiveSize2D(cx, cy))
            off += cy
            total += 1
        base = (R1 - R0 + 1) * ROW_EMU
        assert base >= off, ("土台が足りない", ws.title, R0, base / EMU, off / EMU)
        print("  A%-5d:%s%-5d %3d行 土台%7.1fpx  画像 %dx%-8.1f 余白%6.2fpx  %d枚"
              % (R0, get_column_letter(c1), R1, R1 - R0 + 1, base / EMU,
                 w_px, off / EMU, (base - off) / EMU, len(ch)))

wb.save(dst)
print("\n%d 枚を組み替えた -> %s" % (total, dst))
