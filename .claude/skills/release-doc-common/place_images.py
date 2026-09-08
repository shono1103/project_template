"""relayout.py が作った土台に、画像を**実寸**で貼る。

    python place_images.py <shell.xlsx> <out.xlsx>

`<shell.xlsx>.plan` (relayout.py が書く 5 フィールド TSV
`シート名 ⇥ 土台のアドレス ⇥ 画像パス ⇥ 高さpx ⇥ 幅px`) を読む。

## `oneCellAnchor` + 明示 `<ext>` で貼る

    <xdr:oneCellAnchor>
      <xdr:from><xdr:col>0</..><xdr:colOff>0</..><xdr:row>12</..><xdr:rowOff>0</..></xdr:from>
      <xdr:ext cx="6791325" cy="19488150"/>   <- 画像の実寸 (EMU)

**`twoCellAnchor` は使わない。** `twoCellAnchor` は左上と右下のセルだけを持ち、
`<a:stretch/>` と併せて「その矩形に引き伸ばす」意味になる。矩形の高さは行高の合計から
決まるので、**行高を丸める描画系 (Google スプレッドシート等) では縦横比が崩れる。**
`oneCellAnchor` なら左上だけがセル参照で、大きさは `<ext>` の実寸で決まるため、
どこで開いても歪まない。

`cx` は土台の幅 (plan の幅px)、`cy` は `cx * 画像の高さ / 画像の幅` で出す。
**縦横比は画像そのものから取る**ので、plan の高さpx とずれても画像が正になる。

同じ土台に複数枚 (split_tall.py で割った継ぎ) が来るときは plan の順に積み、
2 枚目以降は直前までの合計 EMU を `row` + `rowOff` に割り振って真下に詰める
(1 行 = 24px = 228600 EMU)。継ぎ目に隙間が出ない。
"""
import io, os, sys
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils import coordinate_to_tuple, get_column_letter

EMU, ROW_PX = 9525, 24
ROW_EMU = ROW_PX * EMU

if len(sys.argv) != 3:
    sys.exit(__doc__)
src, dst = sys.argv[1], sys.argv[2]
wb = load_workbook(src)

ends = {}
for ws in wb.worksheets:
    for rng in ws.merged_cells.ranges:
        ends[(ws.title, rng.min_row, rng.min_col)] = (rng.max_row, rng.max_col)

off = {}                                   # (シート, 土台) -> 積み上げた EMU
n = 0
for line in io.open(src + ".plan", encoding="utf-8").read().splitlines():
    f = line.split("\t")
    if len(f) != 5:
        sys.exit("plan が %d フィールド (期待 5): %s" % (len(f), line[:80]))
    title, addr, path, _h, w_px = f
    ws = wb[title]
    r0, c0 = coordinate_to_tuple(addr)
    if (title, r0, c0) not in ends:
        sys.exit("%s!%s に結合セル (土台) が無い" % (title, addr))
    r1, c1 = ends[(title, r0, c0)]
    ws.cell(r0, c0).value = None                       # プレースホルダを消す

    xi = XLImage(path)
    cx = int(w_px) * EMU
    cy = round(cx * xi.height / xi.width)              # 縦横比は画像の実寸から
    o = off.get((title, addr), 0)
    xi.anchor = OneCellAnchor(
        _from=AnchorMarker(col=c0 - 1, colOff=0,
                           row=r0 - 1 + o // ROW_EMU, rowOff=o % ROW_EMU),
        ext=XDRPositiveSize2D(cx, cy))
    ws.add_image(xi)
    off[(title, addr)] = o + cy
    n += 1
    print("%-22s %-6s %4dx%-6d %5.1fpx から  %s"
          % (title, addr, xi.width, xi.height, o / EMU, os.path.basename(path)))

for (title, addr), used in sorted(off.items()):
    ws = wb[title]
    r0, c0 = coordinate_to_tuple(addr)
    r1, c1 = ends[(title, r0, c0)]
    base = (r1 - r0 + 1) * ROW_EMU
    if base < used:                                    # 土台が画像より低い = ラベル行に被る
        sys.exit("土台が足りない: %s!%s:%s%d  土台 %.1fpx < 画像 %.1fpx"
                 % (title, addr, get_column_letter(c1), r1, base / EMU, used / EMU))

wb.save(dst)
print("\n%d 枚を貼った -> %s  %d bytes" % (n, dst, os.path.getsize(dst)))
